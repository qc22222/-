#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
2. 文本改写器（占位实现）

- 输入：流程化批量视频/2_text_rewriter/input/original_texts/*.docx
- 输出：流程化批量视频/2_text_rewriter/output/rewritten_texts/*.docx

说明：
为便于自动化与联调，本脚本提供一个无依赖的占位改写逻辑：逐段复制原文至新文档，
可选地为每个段落前加上一个“改写：”前缀。后续可替换为 Hugging Face/本地微调模型调用。

环境变量：
  - 可使用 .env 管理（见项目根目录）。本脚本当前不强制需要任何密钥。

示例：
  python 2_text_rewriter/w_to_y.py \
    --input-dir "流程化批量视频/2_text_rewriter/input/original_texts" \
    --output-dir "流程化批量视频/2_text_rewriter/output/rewritten_texts" \
    --prefix
"""

import argparse
import os
from pathlib import Path
from typing import List

from docx import Document
from dotenv import load_dotenv

DEFAULT_INPUT_DIR = Path(__file__).resolve().parent / 'input' / 'original_texts'
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent / 'output' / 'rewritten_texts'


def list_docx_files(input_dir: Path) -> List[Path]:
    return sorted([p for p in input_dir.glob('*.docx') if p.is_file()])


def rewrite_docx(src: Path, dst: Path, add_prefix: bool = False) -> None:
    doc = Document(src)
    out = Document()
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            out.add_paragraph('')
            continue
        new_text = f'改写：{text}' if add_prefix else text
        out.add_paragraph(new_text)
    dst.parent.mkdir(parents=True, exist_ok=True)
    out.save(dst)


def main() -> None:
    load_dotenv()  # 预留未来使用 .env

    parser = argparse.ArgumentParser(description='文本改写器（占位实现）')
    parser.add_argument('--input-dir', type=str, default=str(DEFAULT_INPUT_DIR), help='输入 DOCX 目录')
    parser.add_argument('--output-dir', type=str, default=str(DEFAULT_OUTPUT_DIR), help='输出 DOCX 目录')
    parser.add_argument('--prefix', action='store_true', help='是否为每段添加“改写：”前缀')

    args = parser.parse_args()

    in_dir = Path(args.input_dir)
    out_dir = Path(args.output_dir)

    in_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)

    files = list_docx_files(in_dir)
    if not files:
        print(f'未在 {in_dir} 找到 .docx 文件')
        return

    for src in files:
        dst = out_dir / src.name
        rewrite_docx(src, dst, add_prefix=args.prefix)
        print(f'处理完成: {src.name} -> {dst}')

    print(f'全部完成，输出目录：{out_dir}')


if __name__ == '__main__':
    main()
